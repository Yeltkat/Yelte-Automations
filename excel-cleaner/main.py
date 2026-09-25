"""
Excel/CSV cleaner: merges messy order exports into one clean Excel report.

What it does:
  1. Reads every .csv / .xlsx / .xls file in the input folder
  2. Maps different header names ("E-mail", "Email Address", ...) to one set of columns
  3. Cleans names, emails, phone numbers, dates, prices, quantities and countries
  4. Removes empty rows and duplicates (also across files)
  5. Moves rows that can't be fixed automatically to a "Needs Review" sheet
  6. Writes a formatted Excel report with a summary and a cleaning log

Usage:
  python main.py                       # uses input/ and output/clean_orders_report.xlsx
  python main.py --input-dir exports --output report.xlsx
"""

import argparse
import re
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).parent

# Every known spelling of a header, mapped to the column name we want.
# Headers are compared after lowercasing and removing spaces/punctuation.
COLUMN_ALIASES = {
    "order_id": ["orderid", "ordernumber", "orderno", "order"],
    "customer": ["customername", "customer", "name", "fullname"],
    "email": ["email", "emailaddress", "mail"],
    "phone": ["phone", "telephone", "phonenumber", "tel", "mobile"],
    "order_date": ["orderdate", "date", "created", "createdat"],
    "product": ["product", "item", "productname"],
    "quantity": ["qty", "quantity", "amount"],
    "unit_price": ["unitprice", "price", "priceeach"],
    "country": ["country", "land"],
}

COUNTRY_MAP = {
    "nl": "Netherlands", "netherlands": "Netherlands", "thenetherlands": "Netherlands",
    "nederland": "Netherlands", "holland": "Netherlands",
    "de": "Germany", "germany": "Germany", "deutschland": "Germany", "duitsland": "Germany",
    "be": "Belgium", "belgium": "Belgium", "belgie": "Belgium", "belgië": "Belgium",
}

# Tried in order. Day-first formats come before month-first ones because the
# sample data is European. Adjust for US clients.
DATE_FORMATS = ["%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%b %d %Y", "%d %b %Y"]

DEFAULT_PHONE_PREFIX = "+31"  # used for local numbers that start with a single 0
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[a-z]{2,}$")


# ---------------------------------------------------------------------------
# Reading and merging
# ---------------------------------------------------------------------------

def header_key(header):
    return re.sub(r"[^a-z0-9]", "", str(header).lower())


def standardize_columns(df, source):
    lookup = {alias: target for target, aliases in COLUMN_ALIASES.items() for alias in aliases}
    renamed = {}
    for col in df.columns:
        target = lookup.get(header_key(col))
        if target is None:
            print(f"  ! {source}: unknown column '{col}' is ignored")
            continue
        renamed[col] = target
    df = df.rename(columns=renamed)[list(renamed.values())]
    # Make sure every expected column exists, even if a file is missing one
    for target in COLUMN_ALIASES:
        if target not in df.columns:
            df[target] = ""
    return df[list(COLUMN_ALIASES)]


def read_file(path):
    if path.suffix.lower() == ".csv":
        # sep=None lets pandas detect ',' or ';' automatically
        df = pd.read_csv(path, sep=None, engine="python", dtype=str, keep_default_na=False)
    else:
        df = pd.read_excel(path, dtype=str, keep_default_na=False)
    df = standardize_columns(df, path.name)
    df["source_file"] = path.name
    return df


def load_inputs(input_dir):
    files = sorted(p for p in input_dir.iterdir()
                   if p.suffix.lower() in {".csv", ".xlsx", ".xls"} and not p.name.startswith("~$"))
    if not files:
        raise SystemExit(f"No .csv or .xlsx files found in {input_dir}")
    print(f"Reading {len(files)} file(s) from {input_dir}/")
    frames = []
    for path in files:
        df = read_file(path)
        print(f"  - {path.name}: {len(df)} rows")
        frames.append(df)
    return pd.concat(frames, ignore_index=True)


# ---------------------------------------------------------------------------
# Cleaning helpers (each one takes a single value and returns the clean value)
# ---------------------------------------------------------------------------

def clean_text(value):
    return re.sub(r"\s+", " ", str(value)).strip()


def clean_name(value):
    name = clean_text(value)
    # Title-case, but keep Dutch/German particles like "de" and "van" lowercase
    particles = {"de", "van", "der", "den", "von", "ten", "ter"}
    words = [w.lower() if w.lower() in particles and i > 0 else w.capitalize()
             for i, w in enumerate(name.split(" "))]
    return " ".join(words)


def clean_email(value):
    email = clean_text(value).lower()
    return email if EMAIL_RE.match(email) else None


def clean_phone(value):
    raw = clean_text(value)
    if not raw:
        return None
    digits = re.sub(r"[^\d+]", "", raw)
    if digits.startswith("00"):
        digits = "+" + digits[2:]
    elif digits.startswith("0"):
        digits = DEFAULT_PHONE_PREFIX + digits[1:]
    return digits if re.fullmatch(r"\+\d{9,14}", digits) else None


def clean_date(value):
    text = clean_text(value)
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def clean_price(value):
    text = re.sub(r"(?i)[€$\s]|eur", "", clean_text(value))
    if not text:
        return None
    # Whichever separator comes last is the decimal separator: 1.234,50 or 1,234.50
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    else:
        text = text.replace(",", ".")
    try:
        price = round(float(text), 2)
    except ValueError:
        return None
    return price if price >= 0 else None


def clean_quantity(value):
    try:
        qty = float(clean_text(value))
    except ValueError:
        return None
    return int(qty) if qty > 0 and qty.is_integer() else None


def clean_country(value):
    key = clean_text(value).lower().replace(" ", "")
    return COUNTRY_MAP.get(key, clean_text(value).title() or None)


def clean_product(value):
    return clean_text(value).title()


# ---------------------------------------------------------------------------
# Pipeline
# ---------------------------------------------------------------------------

def clean(df):
    log = [("Rows read from all files", len(df))]

    # 1. Drop completely empty rows
    data_cols = list(COLUMN_ALIASES)
    empty = df[data_cols].apply(lambda col: col.map(clean_text)).eq("").all(axis=1)
    df = df[~empty].copy()
    log.append(("Empty rows removed", int(empty.sum())))

    # 2. Clean every column
    originals = df.copy()
    df["order_id"] = df["order_id"].map(clean_text).str.upper()
    df["customer"] = df["customer"].map(clean_name)
    df["email"] = df["email"].map(clean_email)
    df["phone"] = df["phone"].map(clean_phone)
    df["order_date"] = df["order_date"].map(clean_date)
    df["product"] = df["product"].map(clean_product)
    df["quantity"] = df["quantity"].map(clean_quantity)
    df["unit_price"] = df["unit_price"].map(clean_price)
    df["country"] = df["country"].map(clean_country)

    changed = (originals[data_cols].astype(str) != df[data_cols].astype(str)).sum().sum()
    log.append(("Individual cells reformatted", int(changed)))

    # 3. Remove duplicates. After cleaning, "EMMA JANSEN " and "Emma Jansen" are equal,
    #    so this also catches duplicates that only differed in formatting.
    before = len(df)
    df = df.drop_duplicates(subset=data_cols)
    log.append(("Duplicate rows removed", before - len(df)))

    # When the same order ID appears twice with different data, keep the most complete row
    before = len(df)
    completeness = df[data_cols].notna().sum(axis=1)
    has_id = df["order_id"] != ""
    deduped = (df[has_id].assign(_complete=completeness)
                 .sort_values("_complete", ascending=False, kind="stable")
                 .drop_duplicates(subset="order_id", keep="first")
                 .drop(columns="_complete"))
    df = pd.concat([deduped, df[~has_id]]).sort_index()
    log.append(("Same order ID in multiple rows/files removed", before - len(df)))

    # 4. Split off rows that still have problems, with a reason per row
    required = {
        "order_id": "missing order ID",
        "email": "invalid email",
        "order_date": "unreadable date",
        "quantity": "invalid quantity",
        "unit_price": "invalid price",
    }
    reasons = pd.Series("", index=df.index)
    for col, reason in required.items():
        bad = df[col].isna() | df[col].astype(str).eq("")
        reasons[bad] = reasons[bad] + reason + "; "
    needs_review = df[reasons != ""].copy()
    needs_review["issue"] = reasons[reasons != ""].str.rstrip("; ")
    # Show the original, uncleaned values so a human can fix them
    needs_review = pd.concat(
        [originals.loc[needs_review.index, data_cols + ["source_file"]], needs_review["issue"]],
        axis=1,
    )
    df = df[reasons == ""].copy()
    log.append(("Rows moved to 'Needs Review'", len(needs_review)))

    # 5. Final touches
    df["quantity"] = df["quantity"].astype(int)
    df["total"] = (df["quantity"] * df["unit_price"]).round(2)
    df = df.sort_values(["order_date", "order_id"]).reset_index(drop=True)
    log.append(("Clean rows in final report", len(df)))

    return df, needs_review, log


def build_summary(df):
    months = df["order_date"].dt.to_period("M").astype(str)
    by_month = df.groupby(months)["total"].agg(["count", "sum"]).reset_index()
    by_month.columns = ["Month", "Orders", "Revenue"]

    by_country = df.groupby("country")["total"].agg(["count", "sum"]).reset_index()
    by_country.columns = ["Country", "Orders", "Revenue"]

    by_product = (df.groupby("product")
                    .agg(Units=("quantity", "sum"), Revenue=("total", "sum"))
                    .sort_values("Revenue", ascending=False)
                    .reset_index()
                    .rename(columns={"product": "Product"}))
    return by_month, by_country, by_product


# ---------------------------------------------------------------------------
# Excel output
# ---------------------------------------------------------------------------

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=13)
EURO = '€#,##0.00'


def style_table(ws, start_row=1, start_col=1, money_cols=(), date_cols=()):
    """Style a table whose header is at (start_row, start_col)."""
    for cell in ws[start_row][start_col - 1:]:
        if cell.value is None:
            break
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for row in ws.iter_rows(min_row=start_row + 1):
        for cell in row:
            header = ws.cell(row=start_row, column=cell.column).value
            if header in money_cols:
                cell.number_format = EURO
            elif header in date_cols:
                cell.number_format = "DD-MM-YYYY"


def autofit(ws):
    for col_cells in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(width + 3, 50)


def pretty_headers(df):
    """order_id -> Order ID, unit_price -> Unit Price"""
    return df.rename(columns=lambda c: c.replace("_", " ").title().replace(" Id", " ID"))


def write_report(df, needs_review, log, output_path):
    output_path.parent.mkdir(parents=True, exist_ok=True)
    by_month, by_country, by_product = build_summary(df)


    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        pretty_headers(df).to_excel(writer, sheet_name="Clean Orders", index=False)

        # Summary sheet: three small tables stacked vertically
        row = 3
        for title, table in [("Revenue per month", by_month),
                             ("Revenue per country", by_country),
                             ("Best-selling products", by_product)]:
            table.to_excel(writer, sheet_name="Summary", index=False, startrow=row)
            writer.sheets["Summary"].cell(row=row, column=1, value=title).font = TITLE_FONT
            row += len(table) + 4

        pretty_headers(needs_review).to_excel(writer, sheet_name="Needs Review", index=False)
        pd.DataFrame(log, columns=["Step", "Count"]).to_excel(
            writer, sheet_name="Cleaning Log", index=False)

        # --- formatting ---
        ws = writer.sheets["Clean Orders"]
        style_table(ws, money_cols={"Unit Price", "Total"}, date_cols={"Order Date"})
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        autofit(ws)

        ws = writer.sheets["Summary"]
        ws["A1"] = "Order report"
        ws["A1"].font = Font(bold=True, size=16)
        ws["A2"] = (f"{len(df)} orders  |  total revenue €{df['total'].sum():,.2f}  |  "
                    f"generated {datetime.now():%d-%m-%Y %H:%M}")
        row = 4
        for table in (by_month, by_country, by_product):
            style_table(ws, start_row=row + 1, money_cols={"Revenue"})
            row += len(table) + 4
        autofit(ws)
        ws.column_dimensions["A"].width = 24

        ws = writer.sheets["Needs Review"]
        style_table(ws)
        ws.freeze_panes = "A2"
        autofit(ws)

        ws = writer.sheets["Cleaning Log"]
        style_table(ws)
        autofit(ws)


def main():
    parser = argparse.ArgumentParser(description="Merge and clean messy order exports.")
    parser.add_argument("--input-dir", type=Path, default=BASE_DIR / "input")
    parser.add_argument("--output", type=Path, default=BASE_DIR / "output" / "clean_orders_report.xlsx")
    args = parser.parse_args()

    raw = load_inputs(args.input_dir)
    df, needs_review, log = clean(raw)
    write_report(df, needs_review, log, args.output)

    print("\nCleaning log:")
    for step, count in log:
        print(f"  {step:<48} {count:>5}")
    print(f"\nReport saved to {args.output}")


if __name__ == "__main__":
    main()
