"""
Web scraper: collects every book from books.toscrape.com into CSV and Excel.

books.toscrape.com is a practice site built specifically for learning web
scraping, so it's safe to scrape. The same approach works for public product
catalogs, price lists and directories.

What it does:
  1. Checks robots.txt before fetching anything
  2. Finds all categories from the sidebar
  3. Walks through every page of every category (follows the "next" button)
  4. Optionally opens each book's page for extra details (UPC, stock count, description)
  5. Saves the results as CSV and as a formatted Excel file with a summary sheet

Usage:
  python main.py                          # all categories, listing pages only (~1 minute)
  python main.py --category Travel        # a single category
  python main.py --limit 3 --details      # first 3 categories, including detail pages
"""

import argparse
import re
import time
from datetime import datetime
from pathlib import Path
from urllib.parse import urljoin
from urllib.robotparser import RobotFileParser

import pandas as pd
import requests
from bs4 import BeautifulSoup
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

BASE_URL = "https://books.toscrape.com/"
OUTPUT_DIR = Path(__file__).parent / "output"
USER_AGENT = "Mozilla/5.0 (compatible; portfolio-demo-scraper/1.0)"
REQUEST_DELAY = 0.5  # seconds between requests, to be polite to the server

RATING_WORDS = {"One": 1, "Two": 2, "Three": 3, "Four": 4, "Five": 5}


class Scraper:
    def __init__(self, delay=REQUEST_DELAY):
        self.delay = delay
        self.session = requests.Session()
        self.session.headers["User-Agent"] = USER_AGENT
        # Retry automatically on temporary server errors and rate limiting
        retries = Retry(total=3, backoff_factor=1, status_forcelist=[429, 500, 502, 503, 504])
        self.session.mount("https://", HTTPAdapter(max_retries=retries))
        self.robots = self._load_robots()
        self.requests_made = 0

    def _load_robots(self):
        robots = RobotFileParser()
        response = self.session.get(urljoin(BASE_URL, "robots.txt"), timeout=15)
        # No robots.txt (404) means there are no restrictions
        robots.parse(response.text.splitlines() if response.ok else [])
        return robots

    def get_soup(self, url):
        if not self.robots.can_fetch(USER_AGENT, url):
            raise PermissionError(f"robots.txt does not allow fetching {url}")
        time.sleep(self.delay)
        response = self.session.get(url, timeout=15)
        response.raise_for_status()
        response.encoding = "utf-8"  # the site is UTF-8; without this '£' becomes 'Â£'
        self.requests_made += 1
        return BeautifulSoup(response.text, "html.parser")

    # --- listing pages -----------------------------------------------------

    def get_categories(self):
        soup = self.get_soup(BASE_URL)
        links = soup.select(".side_categories ul li ul li a")
        return {a.get_text(strip=True): urljoin(BASE_URL, a["href"]) for a in links}

    def scrape_category(self, name, url):
        books = []
        page = 1
        while url:
            soup = self.get_soup(url)
            for card in soup.select("article.product_pod"):
                books.append(self.parse_card(card, url, name))
            print(f"  {name:<22} page {page}: {len(books)} books so far")
            next_link = soup.select_one("li.next a")
            url = urljoin(url, next_link["href"]) if next_link else None
            page += 1
        return books

    @staticmethod
    def parse_card(card, page_url, category):
        link = card.select_one("h3 a")
        rating_class = card.select_one("p.star-rating")["class"]
        rating = next((RATING_WORDS[c] for c in rating_class if c in RATING_WORDS), None)
        return {
            "title": link["title"],
            "category": category,
            "price_gbp": parse_price(card.select_one(".price_color").get_text()),
            "rating": rating,
            "in_stock": "In stock" in card.select_one(".availability").get_text(),
            "url": urljoin(page_url, link["href"]),
            "image_url": urljoin(page_url, card.select_one("img")["src"]),
        }

    # --- detail pages ------------------------------------------------------

    def scrape_details(self, book):
        soup = self.get_soup(book["url"])
        table = {row.th.get_text(strip=True): row.td.get_text(strip=True)
                 for row in soup.select("table.table tr")}
        stock = re.search(r"(\d+) available", table.get("Availability", ""))
        description = soup.select_one("#product_description + p")
        return {
            "upc": table.get("UPC"),
            "stock_count": int(stock.group(1)) if stock else 0,
            "reviews": int(table.get("Number of reviews", 0)),
            "description": description.get_text(strip=True) if description else "",
        }


def parse_price(text):
    return float(re.sub(r"[^\d.]", "", text))


def save_results(df, output_dir):
    output_dir.mkdir(parents=True, exist_ok=True)
    csv_path = output_dir / "books.csv"
    xlsx_path = output_dir / "books.xlsx"

    df.to_csv(csv_path, index=False, encoding="utf-8-sig")  # -sig so Excel shows '£' etc. correctly

    summary = (df.groupby("category")
                 .agg(books=("title", "count"),
                      avg_price_gbp=("price_gbp", "mean"),
                      avg_rating=("rating", "mean"),
                      in_stock=("in_stock", "sum"))
                 .round(2)
                 .sort_values("books", ascending=False)
                 .reset_index())

    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Books", index=False)
        summary.to_excel(writer, sheet_name="Per Category", index=False)
        for ws in writer.sheets.values():
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for col in ws.columns:
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[get_column_letter(col[0].column)].width = min(width + 2, 60)
        # Make URLs clickable
        ws = writer.sheets["Books"]
        headers = [c.value for c in ws[1]]
        for name in ("url", "image_url"):
            col = headers.index(name) + 1
            for row in range(2, ws.max_row + 1):
                cell = ws.cell(row=row, column=col)
                cell.hyperlink = cell.value
                cell.style = "Hyperlink"

    return csv_path, xlsx_path


def main():
    parser = argparse.ArgumentParser(description="Scrape books.toscrape.com")
    parser.add_argument("--category", help="scrape only this category, e.g. 'Travel'")
    parser.add_argument("--limit", type=int, help="scrape only the first N categories")
    parser.add_argument("--details", action="store_true",
                        help="also open every book page for UPC, stock count and description (slower)")
    parser.add_argument("--delay", type=float, default=REQUEST_DELAY, help="seconds between requests")
    parser.add_argument("--output-dir", type=Path, default=OUTPUT_DIR)
    args = parser.parse_args()

    started = time.time()
    scraper = Scraper(delay=args.delay)
    categories = scraper.get_categories()
    print(f"Found {len(categories)} categories")

    if args.category:
        match = {k: v for k, v in categories.items() if k.lower() == args.category.lower()}
        if not match:
            raise SystemExit(f"Unknown category '{args.category}'. Options: {', '.join(categories)}")
        categories = match
    if args.limit:
        categories = dict(list(categories.items())[:args.limit])

    books = []
    for name, url in categories.items():
        try:
            books.extend(scraper.scrape_category(name, url))
        except requests.RequestException as exc:
            # One broken category shouldn't stop the whole run
            print(f"  ! skipped {name}: {exc}")

    if args.details:
        print(f"\nFetching details for {len(books)} books...")
        for i, book in enumerate(books, 1):
            try:
                book.update(scraper.scrape_details(book))
            except requests.RequestException as exc:
                print(f"  ! no details for {book['title']}: {exc}")
            if i % 25 == 0 or i == len(books):
                print(f"  {i}/{len(books)}")

    if not books:
        raise SystemExit("No books scraped.")

    df = pd.DataFrame(books).drop_duplicates(subset="url")
    df["scraped_at"] = datetime.now().strftime("%Y-%m-%d %H:%M")
    csv_path, xlsx_path = save_results(df, args.output_dir)

    print(f"\nScraped {len(df)} books from {len(categories)} categories "
          f"with {scraper.requests_made} requests in {time.time() - started:.0f}s")
    print(f"Saved {csv_path}\nSaved {xlsx_path}")


if __name__ == "__main__":
    main()
